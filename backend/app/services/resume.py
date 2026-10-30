"""Resume text extraction and skill detection.

Phase 1 is deterministic: pull the text out of a PDF or DOCX and match it against
a known skill vocabulary (every skill our jobs ask for, plus common aliases). No
AI call, so it is free, fast and predictable. Phase 2 swaps in Claude for richer
parsing (experience, education, titles) and keeps this as the fallback.
"""

from __future__ import annotations

import io
import re

# 4 MB: resumes are rarely over 1 MB, and this stays under the 4.5 MB request cap on
# Vercel Functions in case the /api proxy ever runs through one.
MAX_BYTES = 4 * 1024 * 1024


class ResumeError(ValueError):
    """A problem the user can fix (wrong type, unreadable, too big)."""


# canonical skill -> extra spellings people write on resumes
ALIASES: dict[str, list[str]] = {
    "JavaScript": ["javascript", "js", "es6"],
    "TypeScript": ["typescript", "ts"],
    "Node.js": ["node.js", "nodejs", "node js"],
    "Next.js": ["next.js", "nextjs"],
    "React": ["react", "react.js", "reactjs"],
    "PostgreSQL": ["postgresql", "postgres"],
    "MongoDB": ["mongodb", "mongo"],
    "Machine Learning": ["machine learning", "ml"],
    "Deep Learning": ["deep learning"],
    "LLMs": ["llm", "llms", "large language model", "large language models"],
    "RAG": ["rag", "retrieval augmented generation", "retrieval-augmented generation"],
    "Prompt Engineering": ["prompt engineering"],
    "OpenAI": ["openai", "gpt-4", "chatgpt"],
    "Vector Databases": ["vector database", "vector databases", "pinecone", "faiss", "chroma", "pgvector"],
    "Scikit-learn": ["scikit-learn", "sklearn"],
    "Power BI": ["power bi", "powerbi"],
    "REST APIs": ["rest api", "rest apis", "restful"],
    "CI/CD": ["ci/cd", "ci cd", "github actions", "jenkins"],
    "Data Structures": ["data structures", "dsa"],
    "System Design": ["system design"],
    "AWS": ["aws", "amazon web services"],
    "Azure": ["azure"],
    "Kubernetes": ["kubernetes", "k8s"],
    "API Testing": ["api testing", "postman"],
    "Jetpack Compose": ["jetpack compose"],
    "User Research": ["user research", "usability testing"],
    "Design Systems": ["design system", "design systems"],
    "Stakeholder Management": ["stakeholder management"],
    "Requirements Gathering": ["requirements gathering"],
    "Product Strategy": ["product strategy"],
    "Network Security": ["network security"],
    "Incident Response": ["incident response"],
}

# Skills our jobs use (kept in sync with app.seed.ROLES) plus popular extras.
BASE_SKILLS = [
    "Python", "Java", "SQL", "Git", "React", "TypeScript", "JavaScript", "CSS", "Next.js", "Redux",
    "Testing", "Node.js", "PostgreSQL", "REST APIs", "Docker", "Redis", "Microservices", "MongoDB",
    "AWS", "Azure", "LLMs", "LangChain", "RAG", "PyTorch", "TensorFlow", "Vector Databases", "FastAPI",
    "Machine Learning", "MLOps", "Scikit-learn", "Deep Learning", "Spark", "Prompt Engineering",
    "OpenAI", "Statistics", "Pandas", "Tableau", "Excel", "Power BI", "Airflow", "Kafka", "dbt",
    "Kubernetes", "Terraform", "CI/CD", "Linux", "Networking", "Selenium", "Cypress", "API Testing",
    "Product Strategy", "Roadmapping", "User Research", "Analytics", "Stakeholder Management",
    "Wireframing", "Figma", "Prototyping", "Design Systems", "Requirements Gathering", "Kotlin",
    "Android", "Jetpack Compose", "Network Security", "SIEM", "Incident Response", "Data Structures",
    "System Design", "C++", "C#", "Go", "Django", "Flask", "Spring Boot", "NLP", "Computer Vision",
]


# Two-letter abbreviations are only trusted in capitals ("ML", not "ml" inside a word list).
_SHORT_CAPS = {"ML": "Machine Learning", "JS": "JavaScript", "TS": "TypeScript"}


def _build_patterns() -> list[tuple[str, re.Pattern]]:
    out = []
    for skill in BASE_SKILLS:
        if skill == "Go":
            # "go" is an ordinary English word; only the capitalised name or Golang counts.
            out.append((skill, re.compile(r"(?<![\w.])(?:Go|Golang|golang)(?![\w])")))
            continue
        spellings = {skill.lower(), *(a for a in ALIASES.get(skill, []) if len(a) > 2)}
        alt = "|".join(re.escape(x) for x in sorted(spellings, key=len, reverse=True))
        # Boundaries that also respect skills ending in symbols: C++, C#, Node.js, CI/CD.
        out.append((skill, re.compile(rf"(?<![\w+#.])(?:{alt})(?![\w+#])", re.IGNORECASE)))
    for short, skill in _SHORT_CAPS.items():
        out.append((skill, re.compile(rf"(?<![\w.]){short}(?![\w])")))
    return out


_PATTERNS = _build_patterns()


def detect_skills(text: str) -> list[str]:
    """Skills found in the text, in vocabulary order, each once."""
    found: list[str] = []
    for skill, pat in _PATTERNS:
        if skill not in found and pat.search(text):
            found.append(skill)
    return found


def extract_text(filename: str, data: bytes) -> str:
    if not data:
        raise ResumeError("The file is empty.")
    if len(data) > MAX_BYTES:
        raise ResumeError("The file is larger than 4 MB. Upload a smaller PDF or DOCX.")
    name = filename.lower()
    if data[:4] == b"%PDF":
        text = _pdf_text(data)
    elif data[:2] == b"PK" and name.endswith(".docx"):
        text = _docx_text(data)
    elif name.endswith(".doc"):
        raise ResumeError("Old .doc files aren't supported. Save it as PDF or DOCX and upload again.")
    else:
        raise ResumeError("Upload your resume as a PDF or DOCX file.")
    text = re.sub(r"[ \t]+", " ", text).strip()
    if len(text) < 40:
        raise ResumeError(
            "We couldn't read any text in this file. It may be a scanned image. "
            "Upload a PDF exported from Word or Google Docs, or a DOCX."
        )
    return text


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ResumeError("This PDF is password-protected. Remove the password and upload again.")
        pages = reader.pages[:10]  # resumes are short; cap work on huge files
        return "\n".join((p.extract_text() or "") for p in pages)
    except ResumeError:
        raise
    except (PdfReadError, Exception) as e:  # noqa: BLE001 - any parser failure is a bad file
        raise ResumeError("This PDF couldn't be opened. Try exporting it again.") from e


def _docx_text(data: bytes) -> str:
    import docx

    try:
        d = docx.Document(io.BytesIO(data))
    except Exception as e:  # noqa: BLE001
        raise ResumeError("This DOCX couldn't be opened. Try saving it again.") from e
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.extend(c.text for c in row.cells)
    return "\n".join(parts)
