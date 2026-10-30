"""Resume Builder: the resume data model, ATS checks, and PDF/DOCX export.

All three templates are single-column with real text and standard section names,
because that is what applicant tracking systems parse reliably. Templates differ in
typography, colour and header layout, not in structure.
"""

from __future__ import annotations

import io
import re
from collections import Counter
from html import escape

from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Job, User

TEMPLATES = ("modern", "professional", "minimal")


# ---------------------------------------------------------------- data model
class Basics(BaseModel):
    name: str = Field(default="", max_length=120)
    title: str = Field(default="", max_length=120)
    email: str = Field(default="", max_length=200)
    phone: str = Field(default="", max_length=40)
    location: str = Field(default="", max_length=120)
    linkedin: str = Field(default="", max_length=300)
    website: str = Field(default="", max_length=300)
    summary: str = Field(default="", max_length=1500)


class Experience(BaseModel):
    role: str = Field(default="", max_length=120)
    company: str = Field(default="", max_length=120)
    location: str = Field(default="", max_length=120)
    start: str = Field(default="", max_length=30)
    end: str = Field(default="", max_length=30)  # "" or "Present"
    bullets: list[str] = Field(default_factory=list, max_length=10)


class Education(BaseModel):
    degree: str = Field(default="", max_length=160)
    school: str = Field(default="", max_length=160)
    start: str = Field(default="", max_length=30)
    end: str = Field(default="", max_length=30)
    score: str = Field(default="", max_length=40)


class Project(BaseModel):
    name: str = Field(default="", max_length=120)
    link: str = Field(default="", max_length=300)
    bullets: list[str] = Field(default_factory=list, max_length=8)


class ResumeData(BaseModel):
    basics: Basics = Field(default_factory=Basics)
    experience: list[Experience] = Field(default_factory=list, max_length=12)
    projects: list[Project] = Field(default_factory=list, max_length=10)
    education: list[Education] = Field(default_factory=list, max_length=6)
    skills: list[str] = Field(default_factory=list, max_length=60)
    certifications: list[str] = Field(default_factory=list, max_length=15)


def from_profile(user: User) -> ResumeData:
    return ResumeData(
        basics=Basics(name=user.name, title=user.target_role or "", email=user.email, location=user.location or ""),
        skills=list(user.skills or [])[:30],
    )


# ---------------------------------------------------------------- ATS checks
ACTION_VERBS = set(
    """achieved analysed analyzed architected automated boosted built championed coached collaborated conducted
    configured created cut debugged decreased defined delivered deployed designed developed directed drove eliminated
    enabled engineered established evaluated expanded generated grew guided handled identified implemented improved
    increased integrated introduced launched led maintained managed mentored migrated modelled modeled monitored
    negotiated optimised optimized organised organized owned partnered planned presented prioritised prioritized
    produced programmed proposed prototyped published raised rebuilt reduced redesigned refactored resolved restructured
    reviewed revamped saved scaled secured shipped simplified solved spearheaded standardised standardized streamlined
    supported taught tested trained transformed tripled troubleshot upgraded wrote ran executed coordinated contributed
    facilitated initiated forecasted measured mapped visualised visualized translated audited assessed onboarded
    recruited sourced closed won converted calculated researched surveyed documented drafted edited curated
    administered operated installed served assisted volunteered hosted organised founded co-founded""".split()
)
_NUMBER = re.compile(r"\d")


def _bullets(d: ResumeData) -> list[str]:
    return [b.strip() for e in d.experience for b in e.bullets if b.strip()] + [
        b.strip() for p in d.projects for b in p.bullets if b.strip()
    ]


def role_keywords(db: Session, target_role: str | None, limit: int = 12) -> list[str]:
    """The skills most often asked for in open jobs that match the target role's title."""
    if not target_role:
        return []
    words = [w for w in re.split(r"\W+", target_role.lower()) if len(w) > 2]
    if not words:
        return []
    q = select(Job.skills).where(Job.active.is_(True))
    for w in words:
        q = q.where(Job.title.ilike(f"%{w}%"))
    counts: Counter[str] = Counter()
    for skills in db.scalars(q.limit(400)):
        counts.update(skills or [])
    return [s for s, _ in counts.most_common(limit)]


def ats_report(d: ResumeData, keywords: list[str]) -> dict:
    b = d.basics
    bullets = _bullets(d)
    text = " ".join([b.summary, *bullets, *d.skills, *(e.role for e in d.experience), *(p.name for p in d.projects)]).lower()
    summary_words = len(b.summary.split())
    first_words = [re.sub(r"[^a-z]", "", x.split()[0].lower()) for x in bullets if x.split()]
    verb_share = sum(w in ACTION_VERBS for w in first_words) / len(first_words) if first_words else 0
    number_share = sum(bool(_NUMBER.search(x)) for x in bullets) / len(bullets) if bullets else 0
    good_length = sum(8 <= len(x.split()) <= 32 for x in bullets) / len(bullets) if bullets else 0
    found = [k for k in keywords if re.search(rf"(?<![\w+#.]){re.escape(k.lower())}(?![\w+#])", text)]
    missing = [k for k in keywords if k not in found]
    coverage = len(found) / len(keywords) if keywords else None

    checks = [
        ("contact", 10, bool(b.email and b.phone), "Email and phone number are in the header.",
         "Add your email and phone number so recruiters can reach you."),
        ("links", 5, bool(b.linkedin or b.website), "Includes a LinkedIn or portfolio link.",
         "Add your LinkedIn URL (or GitHub/portfolio). Many recruiters check it first."),
        ("summary", 10, 25 <= summary_words <= 90, f"Summary is {summary_words} words: the right length.",
         f"Write a 2–4 sentence summary (25–90 words) aimed at your target role. It's {summary_words} words now."),
        ("experience", 15, bool(d.experience or d.projects), "Has experience or projects to show your work.",
         "Add at least one job, internship or project. Freshers: projects count."),
        ("verbs", 10, bool(bullets) and verb_share >= 0.7, f"{round(verb_share * 100)}% of bullets start with an action verb.",
         f"Start bullets with an action verb (Built, Led, Reduced…). {round(verb_share * 100)}% do now; aim for 70%+."),
        ("numbers", 15, bool(bullets) and number_share >= 0.4, f"{round(number_share * 100)}% of bullets show a measurable result.",
         f"Add numbers to more bullets (users, %, time saved, ₹). {round(number_share * 100)}% have one; aim for 40%+."),
        ("length", 5, bool(bullets) and good_length >= 0.8, "Bullets are a readable length.",
         "Keep each bullet to one or two lines (8–32 words)."),
        ("education", 10, bool(d.education), "Education is listed.", "Add your degree, college and year."),
        ("skills", 5, len(d.skills) >= 6, f"{len(d.skills)} skills listed.", f"List at least 6 relevant skills. You have {len(d.skills)}."),
    ]
    if coverage is None:
        checks.append(("keywords", 15, False, "", "Set a target role in your profile to check keywords recruiters search for."))
    else:
        checks.append(("keywords", 15, coverage >= 0.6, f"Covers {len(found)} of {len(keywords)} keywords jobs for your role ask for.",
                       f"Covers {len(found)} of {len(keywords)} keywords jobs for your role ask for. Add the ones you genuinely have."))

    out, score = [], 0
    for cid, weight, ok, good, bad in checks:
        score += weight if ok else 0
        out.append({"id": cid, "weight": weight, "pass": ok, "message": good if ok else bad})
    return {"score": score, "checks": out, "keywords_found": found, "keywords_missing": missing}


# ---------------------------------------------------------------- export
_STYLE = {
    "modern": {"accent": "#5B4BD6", "font": "sans", "center": False, "rule": True},
    "professional": {"accent": "#16183A", "font": "serif", "center": True, "rule": True},
    "minimal": {"accent": "#16183A", "font": "sans", "center": False, "rule": False},
}


def _dates(a: str, b: str) -> str:
    return " – ".join(x for x in (a.strip(), b.strip()) if x)


def _contact_line(b: Basics) -> list[str]:
    return [x for x in (b.email, b.phone, b.location, b.linkedin, b.website) if x.strip()]


def _fonts() -> tuple[str, str, str, str]:
    """Unicode TTFs when the system has DejaVu (Docker image installs it), else built-ins."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    base = "/usr/share/fonts/truetype/dejavu/"
    try:
        for name, f in (("DV", "DejaVuSans.ttf"), ("DV-B", "DejaVuSans-Bold.ttf"), ("DVS", "DejaVuSerif.ttf"), ("DVS-B", "DejaVuSerif-Bold.ttf")):
            if name not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont(name, base + f))
        return "DV", "DV-B", "DVS", "DVS-B"
    except Exception:  # noqa: BLE001 - fonts missing: fall back to the PDF standard fonts
        return "Helvetica", "Helvetica-Bold", "Times-Roman", "Times-Bold"


def to_pdf(d: ResumeData, template: str) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import HRFlowable, KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    st = _STYLE.get(template, _STYLE["modern"])
    sans, sans_b, serif, serif_b = _fonts()
    reg, bold = (serif, serif_b) if st["font"] == "serif" else (sans, sans_b)
    accent = colors.HexColor(st["accent"])
    ink, soft = colors.HexColor("#16183A"), colors.HexColor("#4B4F72")
    align = TA_CENTER if st["center"] else TA_LEFT

    S = {
        "name": ParagraphStyle("n", fontName=bold, fontSize=20, leading=24, textColor=accent if template == "modern" else ink, alignment=align),
        "title": ParagraphStyle("t", fontName=reg, fontSize=11, leading=14, textColor=soft, alignment=align),
        "contact": ParagraphStyle("c", fontName=reg, fontSize=9, leading=12, textColor=soft, alignment=align, spaceBefore=3),
        "h": ParagraphStyle("h", fontName=bold, fontSize=11, leading=14, textColor=accent, spaceBefore=10, spaceAfter=2),
        "body": ParagraphStyle("b", fontName=reg, fontSize=9.5, leading=13, textColor=ink),
        "item": ParagraphStyle("i", fontName=bold, fontSize=10, leading=13, textColor=ink),
        "meta": ParagraphStyle("m", fontName=reg, fontSize=9, leading=12, textColor=soft),
        "bullet": ParagraphStyle("bl", fontName=reg, fontSize=9.5, leading=13, textColor=ink, leftIndent=10, bulletIndent=0),
    }
    e = lambda t: escape(t or "")  # noqa: E731
    flow = []
    b = d.basics
    flow.append(Paragraph(e(b.name) or "Your Name", S["name"]))
    if b.title:
        flow.append(Paragraph(e(b.title), S["title"]))
    if _contact_line(b):
        flow.append(Paragraph("  |  ".join(e(x) for x in _contact_line(b)), S["contact"]))

    def heading(text: str):
        flow.append(Paragraph(text.upper() if template != "minimal" else text, S["h"]))
        if st["rule"]:
            flow.append(HRFlowable(width="100%", thickness=0.6, color=accent, spaceBefore=0, spaceAfter=4))

    def row(left: str, right: str):
        t = Table([[Paragraph(left, S["item"]), Paragraph(e(right), ParagraphStyle("r", parent=S["meta"], alignment=2))]], colWidths=["72%", "28%"])
        t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 0), ("TOPPADDING", (0, 0), (-1, -1), 0)]))
        return t

    if b.summary.strip():
        heading("Summary")
        flow.append(Paragraph(e(b.summary), S["body"]))
    if d.experience:
        heading("Experience")
        for x in d.experience:
            block = [row(f"{e(x.role)}{' — ' + e(x.company) if x.company else ''}", _dates(x.start, x.end))]
            if x.location:
                block.append(Paragraph(e(x.location), S["meta"]))
            block += [Paragraph(e(bl), S["bullet"], bulletText="•") for bl in x.bullets if bl.strip()]
            block.append(Spacer(1, 5))
            flow.append(KeepTogether(block))
    if d.projects:
        heading("Projects")
        for p in d.projects:
            block = [row(e(p.name), p.link)]
            block += [Paragraph(e(bl), S["bullet"], bulletText="•") for bl in p.bullets if bl.strip()]
            block.append(Spacer(1, 5))
            flow.append(KeepTogether(block))
    if d.education:
        heading("Education")
        for ed in d.education:
            flow.append(row(f"{e(ed.degree)}{', ' + e(ed.school) if ed.school else ''}", _dates(ed.start, ed.end)))
            if ed.score:
                flow.append(Paragraph(e(ed.score), S["meta"]))
            flow.append(Spacer(1, 4))
    if d.skills:
        heading("Skills")
        flow.append(Paragraph(e(", ".join(d.skills)), S["body"]))
    if d.certifications:
        heading("Certifications")
        flow += [Paragraph(e(c), S["bullet"], bulletText="•") for c in d.certifications if c.strip()]

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm, topMargin=14 * mm, bottomMargin=14 * mm,
                            title=f"{b.name or 'Resume'} - Resume", author=b.name or "")
    doc.build(flow)
    return buf.getvalue()


def to_docx(d: ResumeData, template: str) -> bytes:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    st = _STYLE.get(template, _STYLE["modern"])
    accent = RGBColor.from_string(st["accent"].lstrip("#"))
    soft = RGBColor(0x4B, 0x4F, 0x72)
    font = "Georgia" if st["font"] == "serif" else "Calibri"

    doc = docx.Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(16)
    sec.top_margin = sec.bottom_margin = Mm(14)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = font, Pt(10)
    normal.paragraph_format.space_after = Pt(2)

    def para(text="", size=10, bold=False, color=None, center=False, after=2):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(after)
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if text:
            r = p.add_run(text)
            r.bold, r.font.size, r.font.name = bold, Pt(size), font
            if color is not None:
                r.font.color.rgb = color
        return p

    def rule(p):
        pPr = p._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", st["accent"].lstrip("#"))):
            bottom.set(qn(k), v)
        bdr.append(bottom)
        pPr.append(bdr)

    def heading(text):
        p = para(text.upper() if template != "minimal" else text, 11, True, accent, after=3)
        p.paragraph_format.space_before = Pt(8)
        if st["rule"]:
            rule(p)

    def item(left, right):
        p = para(after=0)
        r = p.add_run(left)
        r.bold, r.font.name = True, font
        if right:
            p.add_run("\t" + right).font.color.rgb = soft
            p.paragraph_format.tab_stops.add_tab_stop(Mm(178), alignment=2)
        return p

    def bullet(text):
        p = doc.add_paragraph(text, style="List Bullet")
        p.paragraph_format.space_after = Pt(1)

    b = d.basics
    para(b.name or "Your Name", 20, True, accent if template == "modern" else None, st["center"], after=0)
    if b.title:
        para(b.title, 11, False, soft, st["center"], after=0)
    if _contact_line(b):
        para("  |  ".join(_contact_line(b)), 9, False, soft, st["center"], after=4)
    if b.summary.strip():
        heading("Summary")
        para(b.summary)
    if d.experience:
        heading("Experience")
        for x in d.experience:
            item(f"{x.role}{' — ' + x.company if x.company else ''}", _dates(x.start, x.end))
            if x.location:
                para(x.location, 9, color=soft, after=1)
            for bl in x.bullets:
                if bl.strip():
                    bullet(bl)
    if d.projects:
        heading("Projects")
        for p in d.projects:
            item(p.name, p.link)
            for bl in p.bullets:
                if bl.strip():
                    bullet(bl)
    if d.education:
        heading("Education")
        for ed in d.education:
            item(f"{ed.degree}{', ' + ed.school if ed.school else ''}", _dates(ed.start, ed.end))
            if ed.score:
                para(ed.score, 9, color=soft)
    if d.skills:
        heading("Skills")
        para(", ".join(d.skills))
    if d.certifications:
        heading("Certifications")
        for c in d.certifications:
            if c.strip():
                bullet(c)

    doc.core_properties.title = f"{b.name or 'Resume'} - Resume"
    doc.core_properties.author = b.name or ""
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
