import io
import uuid

import pytest

from app.services.resume import detect_skills

RESUME_TEXT = (
    "Neha Verma, Data Analyst. Pune, India. "
    "Built dashboards in Power BI and Tableau. Wrote SQL and Python (Pandas) for weekly reports. "
    "Advanced Excel. Presented findings to stakeholders. Ready to go to market with new KPIs."
)


def _pdf(text: str) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    y = 800
    for line in [text[i : i + 90] for i in range(0, len(text), 90)]:
        c.drawString(40, y, line)
        y -= 16
    c.save()
    return buf.getvalue()


def _docx(text: str) -> bytes:
    import docx

    d = docx.Document()
    d.add_paragraph(text)
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "Tools"
    t.rows[0].cells[1].text = "Airflow, dbt"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _signup(client):
    client.post(
        "/api/auth/signup",
        json={"name": "Neha Verma", "email": f"n{uuid.uuid4().hex[:8]}@example.com", "password": "secret123", "role_type": "fresher"},
    )
    client.patch("/api/me", json={"skills": ["SQL"]})


def test_detect_skills_handles_symbols_aliases_and_false_friends():
    found = detect_skills("C++, C#, Node.js, CI/CD, scikit-learn, ML. Let's go to market. Excellent comms.")
    assert {"C++", "C#", "Node.js", "CI/CD", "Scikit-learn", "Machine Learning"} <= set(found)
    assert "Go" not in found and "Excel" not in found


def test_pdf_upload_adds_new_skills_only(client):
    _signup(client)
    r = client.post("/api/me/resume", files={"file": ("neha.pdf", _pdf(RESUME_TEXT), "application/pdf")})
    assert r.status_code == 200, r.text
    body = r.json()
    assert {"Power BI", "Tableau", "Python", "Pandas", "Excel", "SQL"} <= set(body["detected_skills"])
    assert "SQL" not in body["added_skills"]  # already on the profile
    assert "Go" not in body["detected_skills"]
    assert set(body["user"]["skills"]) >= {"SQL", "Power BI", "Python"}
    assert body["word_count"] > 20


def test_docx_upload_reads_tables_and_replaces_previous(client):
    _signup(client)
    client.post("/api/me/resume", files={"file": ("a.pdf", _pdf(RESUME_TEXT), "application/pdf")})
    r = client.post(
        "/api/me/resume",
        files={"file": ("cv.docx", _docx(RESUME_TEXT), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert r.status_code == 200, r.text
    assert {"Airflow", "dbt"} <= set(r.json()["detected_skills"])
    assert r.json()["filename"] == "cv.docx"


@pytest.mark.parametrize(
    "name,data,needle",
    [
        ("photo.png", b"\x89PNG\r\n\x1a\n" + b"0" * 100, "PDF or DOCX"),
        ("old.doc", b"\xd0\xcf\x11\xe0" + b"0" * 100, ".doc files"),
        ("empty.pdf", b"", "empty"),
        ("broken.pdf", b"%PDF-1.4 not really a pdf", "couldn't be opened"),
        ("huge.pdf", b"%PDF" + b"0" * (4 * 1024 * 1024 + 10), "4 MB"),
    ],
    ids=["png", "doc", "empty", "broken", "too-big"],
)
def test_bad_files_get_clear_errors(client, name, data, needle):
    _signup(client)
    r = client.post("/api/me/resume", files={"file": (name, data, "application/octet-stream")})
    assert r.status_code == 422
    assert needle in r.json()["detail"]


def test_scanned_pdf_with_no_text(client):
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.rect(50, 50, 200, 200, fill=1)  # a shape, no text: like a scanned image
    c.save()
    _signup(client)
    r = client.post("/api/me/resume", files={"file": ("scan.pdf", buf.getvalue(), "application/pdf")})
    assert r.status_code == 422 and "scanned" in r.json()["detail"]


def test_resume_requires_login(client):
    client.cookies.clear()
    r = client.post("/api/me/resume", files={"file": ("a.pdf", _pdf(RESUME_TEXT), "application/pdf")})
    assert r.status_code == 401
